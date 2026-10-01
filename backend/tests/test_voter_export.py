"""WP-12: IT admin voter-register export. Invariants I1-I12 (see playbook 12.4.8)."""
import csv, io
import pytest
from openpyxl import load_workbook

import main
from auth import create_access_token
from tests.test_flows import env, START, Clock  # noqa: F401  (fixture)

pytestmark = pytest.mark.asyncio

FULL_PHONE = "256700111222"


async def seed(e):
    """Voters with data that MUST and MUST NOT appear, including hostile values."""
    await e.db.settings.insert_one({"name": "voter_fields", "org_id": e.org_id, "fields": [
        {"key": "hostel", "label": "Hostel", "enabled": True},
        {"key": "secretnote", "label": "Secret Note", "enabled": False}]})
    await e.voter("23/u/001", "Alpha One", (FULL_PHONE, "256700333444"), has_voted=True, last_status="authenticated",
                  attrs={"hostel": "Block A", "secretnote": "DISABLED-FIELD-VALUE"},
                  otp_hash="OTPHASHVALUE", vote_jti="JTIVALUE", it_admin_password_hash="PWHASHVALUE", added_by="zzz-added-by")
    await e.voter("23/u/002", "=HYPERLINK(\"http://evil\",\"x\")", ("256700555666",), attrs={"hostel": "@cmd"})


async def set_mode(e, mode):
    await e.db.voters.update_one({"student_id": "it1", "org_id": e.org_id}, {"$set": {"it_admin_export_mode": mode}})


async def export(e, headers=None, fmt="csv"):
    return await e.client.post("/admin/voters/export", json={"format": fmt}, headers=headers or e.it)


def csv_rows(resp):
    return list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))


# I1 -----------------------------------------------------------------------------------------------
async def test_T1_default_is_off_and_refused(env):
    await seed(env)
    assert (await env.client.get("/admin/voters/export/permission", headers=env.it)).json()["mode"] == "none"
    r = await export(env)
    assert r.status_code == 403 and FULL_PHONE not in r.text and "Alpha" not in r.text
    a = await env.db.audit_log.find_one({"action": "voter_register_export_denied"})
    assert a and a["details"] == {"reason": "mode_none"}


@pytest.mark.parametrize("bad", ["FULL", "Full", " full", "yes", "", None, ["full"], 1, True, "none"])
async def test_T2_corrupt_mode_fails_closed(env, bad):
    await seed(env); await set_mode(env, bad)
    assert main.normalize_export_mode(bad) == "none"
    assert (await export(env)).status_code == 403


# I5 + D3 ------------------------------------------------------------------------------------------
async def test_T3_redacted_hides_phones_and_equals_on_screen_list(env):
    await seed(env); await set_mode(env, "redacted")
    r = await export(env)
    assert r.status_code == 200
    text = r.content.decode("utf-8-sig")
    assert FULL_PHONE not in text and "256700333444" not in text and "256700555666" not in text
    rows = csv_rows(r)
    assert rows[0][:4] == ["Registration Number", "Full Name", "Phone 1", "Phone 2"]
    listed = (await env.client.get("/admin/voters/list", headers=env.it, params={"page_size": 50})).json()["results"]
    by_id = {v["student_id"].upper(): v for v in listed}
    for row in rows[1:]:
        assert row[2:4][:len(by_id[row[0]]["phone_numbers"])] == by_id[row[0]]["phone_numbers"]   # export == screen


async def test_T3b_mask_function_pinned():
    # `_mask_phone` is defined twice in main.py; the LATER one is effective. If someone removes the duplicate this
    # fails on purpose: confirm the new masking is acceptable, then update this line.
    assert main._mask_phone(FULL_PHONE) == "*********222"


# I6 + I7 ------------------------------------------------------------------------------------------
@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
@pytest.mark.parametrize("mode", ["redacted", "full"])
async def test_T4_T5_allowlist_no_leaks_in_any_mode_or_format(env, mode, fmt):
    await seed(env); await set_mode(env, mode)
    r = await export(env, fmt=fmt)
    assert r.status_code == 200
    blob = r.content.decode("utf-8-sig", "ignore") if fmt == "csv" else " ".join(
        str(c.value) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows() for c in row)
    for forbidden in ["has_voted", "last_status", "authenticated", "OTPHASHVALUE", "JTIVALUE", "PWHASHVALUE",
                      "zzz-added-by", "DISABLED-FIELD-VALUE", "Secret Note", "is_it_admin", "it_admin_export_mode"]:
        assert forbidden not in blob, forbidden
    assert "Block A" in blob and "Hostel" in blob            # enabled optional field IS included


async def test_T4_full_contains_every_stored_value(env):
    await seed(env); await set_mode(env, "full")
    rows = csv_rows(await export(env))
    alpha = next(r for r in rows if r[0] == "23/U/001")
    assert alpha[1] == "Alpha One" and alpha[2] == FULL_PHONE and alpha[3] == "256700333444"
    assert len(rows) - 1 == await env.db.voters.count_documents({"org_id": env.org_id})   # no row dropped


# I2 / I3 ------------------------------------------------------------------------------------------
async def test_T6_only_it_admin_role(env):
    await seed(env); await set_mode(env, "full")
    for hdr in (env.sa, env.com1, env.over, env.tok("fc1", "financial_controller")):
        assert (await export(env, headers=hdr)).status_code == 403
        assert (await env.client.get("/admin/voters/export/permission", headers=hdr)).status_code == 403
    assert (await env.client.post("/admin/voters/export", json={"format": "csv"}, headers={"X-Org-Slug": "t1"})).status_code in (401, 403)  # no token


async def test_T7_view_as_token_cannot_export(env):
    await seed(env); await set_mode(env, "full")
    tok = create_access_token(subject="it1", role="it_admin", org_id=env.org_id, extra_claims={"view_only": True})
    h = {"Authorization": f"Bearer {tok}", "X-Org-Slug": "t1"}
    assert (await env.client.get("/admin/voters/export/permission", headers=h)).status_code == 200
    r = await export(env, headers=h)
    assert r.status_code == 403 and "Read-only" in r.json()["detail"]
    assert await env.db.audit_log.count_documents({"action": "voter_register_export"}) == 0


# I4 -----------------------------------------------------------------------------------------------
async def test_T8_downgrade_applies_to_existing_session(env):
    await seed(env); await set_mode(env, "full")
    assert (await export(env)).status_code == 200                      # same token ...
    r = await env.client.put("/superadmin/it-admins/it1/export-mode", json={"mode": "none"}, headers=env.sa)
    assert r.status_code == 200 and r.json()["changed"] is True
    assert (await export(env)).status_code == 403                      # ... refused immediately, no re-login


# Superadmin control ------------------------------------------------------------------------------
async def test_T9_superadmin_setter_rules_and_audit(env):
    put = lambda mode, hdr=None, sid="it1": env.client.put(f"/superadmin/it-admins/{sid}/export-mode", json={"mode": mode}, headers=hdr or env.sa)
    assert (await put("everything")).status_code == 400
    assert (await put("full", sid="v1")).status_code == 404            # not an IT admin
    assert (await put("full", hdr=env.it)).status_code == 403          # IT admin cannot grant themselves
    assert (await put("full", hdr=env.com1)).status_code == 403
    assert (await put("full")).json()["changed"] is True
    assert (await put("full")).json()["changed"] is False             # idempotent, not re-logged
    logs = [a async for a in env.db.audit_log.find({"action": "it_admin_export_mode_changed"})]
    assert len(logs) == 1 and logs[0]["details"] == {"student_id": "it1", "old": "none", "new": "full"}
    doc = await env.db.voters.find_one({"student_id": "it1"})
    assert doc["it_admin_export_mode"] == "full" and doc["it_admin_export_mode_set_by"]


async def test_T10_toggle_resets_mode_both_directions(env):
    await set_mode(env, "full")
    await env.client.post("/superadmin/it-admins/it1/toggle", headers=env.sa)   # revoke
    await env.client.post("/superadmin/it-admins/it1/toggle", headers=env.sa)   # re-grant
    doc = await env.db.voters.find_one({"student_id": "it1"})
    assert doc["is_it_admin"] is True and "it_admin_export_mode" not in doc


async def test_T11_list_reports_mode_normalised(env):
    await set_mode(env, "WEIRD")
    row = next(a for a in (await env.client.get("/superadmin/it-admins", headers=env.sa)).json() if a["student_id"] == "it1")
    assert row["it_admin_export_mode"] == "none"
    await set_mode(env, "redacted")
    row = next(a for a in (await env.client.get("/superadmin/it-admins", headers=env.sa)).json() if a["student_id"] == "it1")
    assert row["it_admin_export_mode"] == "redacted"


async def test_T11b_commissioner_list_has_no_export_field(env):
    # regression: a stray line once injected it_admin_export_mode into the commissioner roster
    r = await env.client.get("/admin/commissioners", headers=env.com1)
    assert r.status_code == 200 and "it_admin_export_mode" not in r.text


# I8 -----------------------------------------------------------------------------------------------
async def test_T12_tenant_isolation(env):
    await seed(env); await set_mode(env, "full")
    org2 = await env.db.organizations.insert_one({"slug": "t2", "name": "T2"})
    await env.db.voters.insert_one({"student_id": "99/u/999", "full_name": "Other Org Person", "phone_numbers": ["256799999999"],
                                    "org_id": str(org2.inserted_id)})
    text = (await export(env)).content.decode("utf-8-sig")
    assert "Other Org Person" not in text and "256799999999" not in text
    cross = await env.client.post("/admin/voters/export", json={"format": "csv"},
                                  headers={"Authorization": env.it["Authorization"], "X-Org-Slug": "t2"})
    assert cross.status_code == 403                                      # token from org A on org B


# I9 -----------------------------------------------------------------------------------------------
async def test_T13_formula_injection_neutralised(env):
    await seed(env); await set_mode(env, "full")
    evil = next(r for r in csv_rows(await export(env)) if r[0] == "23/U/002")
    assert evil[1].startswith("'=") and evil[-1] == "'@cmd"
    ws = load_workbook(io.BytesIO((await export(env, fmt="xlsx")).content)).active
    cells = {c.value: c for row in ws.iter_rows() for c in row if c.value}
    c = cells['=HYPERLINK("http://evil","x")']
    assert c.data_type == "s"                                            # stored as text, NOT a formula ('f')


async def test_T14_xlsx_roundtrip_phones_are_text(env):
    await seed(env); await set_mode(env, "full")
    r = await export(env, fmt="xlsx")
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    alpha = next(x for x in rows if x[0] == "23/U/001")
    assert alpha[2] == FULL_PHONE and isinstance(alpha[2], str)


# I11 ----------------------------------------------------------------------------------------------
async def test_T15_row_cap(env, monkeypatch):
    await seed(env); await set_mode(env, "full")
    monkeypatch.setattr(main, "EXPORT_MAX_ROWS", 1)
    r = await export(env)
    assert r.status_code == 413
    assert (await env.db.audit_log.find_one({"action": "voter_register_export_denied"}))["details"] == {"reason": "too_many_rows"}


async def test_T16_rate_limit(env, monkeypatch):
    await seed(env); await set_mode(env, "full")
    monkeypatch.setattr(main, "EXPORT_RATE_LIMIT", 2)
    assert [(await export(env)).status_code for _ in range(3)] == [200, 200, 429]


# I10 ----------------------------------------------------------------------------------------------
async def test_T17_audit_has_no_pii(env):
    await seed(env); await set_mode(env, "full")
    await export(env); await export(env, fmt="xlsx")
    logs = [a async for a in env.db.audit_log.find({"action": "voter_register_export"})]
    assert len(logs) == 2 and all(set(a["details"]) == {"mode", "format", "rows"} for a in logs)
    dump = str([{k: v for k, v in a.items() if k != "_id"} async for a in env.db.audit_log.find({})])
    for pii in (FULL_PHONE, "Alpha One", "23/u/001", "23/U/001", "Block A"):
        assert pii not in dump


async def test_T17b_no_data_if_audit_write_fails(env, monkeypatch):
    await seed(env); await set_mode(env, "full")
    real = main.log_action
    async def boom(action, *a, **k):
        if action == "voter_register_export":
            raise RuntimeError("db down")
        return await real(action, *a, **k)
    monkeypatch.setattr(main, "log_action", boom)
    with pytest.raises(Exception):                                       # RuntimeError, or an ExceptionGroup wrapping it
        await export(env)                                                # the point: no 200 response carrying data


async def test_T18_headers_and_cors_expose(env):
    await seed(env); await set_mode(env, "redacted")
    h = {**env.it, "Origin": "http://localhost:5173"}                    # default allowed origin in tests
    r = await export(env, headers=h)
    assert r.headers["cache-control"] == "no-store"
    assert r.headers["x-content-type-options"] == "nosniff"
    assert r.headers["content-disposition"].startswith('attachment; filename="voter-register-t1-redacted-')
    assert r.headers["content-disposition"].endswith('.csv"')
    assert r.headers["x-export-mode"] == "redacted"
    assert r.headers["x-export-rows"] == str(await env.db.voters.count_documents({"org_id": env.org_id}))
    exposed = r.headers.get("access-control-expose-headers", "").lower()
    assert "content-disposition" in exposed and "x-export-rows" in exposed
    assert r.content.startswith(b"\xef\xbb\xbf")                         # UTF-8 BOM for Excel


async def test_T18b_bad_format(env):
    await set_mode(env, "full")
    assert (await export(env, fmt="pdf")).status_code == 400


# I12 + regression ---------------------------------------------------------------------------------
async def test_T19_field_never_leaks_and_existing_endpoints_unchanged(env):
    await seed(env); await set_mode(env, "full")
    lst = await env.client.get("/admin/voters/list", headers=env.it)
    old = await env.client.get("/admin/voters", headers=env.it)
    pub = await env.client.get("/voter-register")
    for r in (lst, old, pub):
        assert r.status_code == 200 and "it_admin_export_mode" not in r.text
    for v in lst.json()["results"]:
        assert "has_voted" not in v and "last_status" not in v                     # IT admin still cannot see voting status
        assert all("*" in p for p in v["phone_numbers"])                            # list still masked for IT admin
    assert FULL_PHONE not in lst.text and FULL_PHONE not in old.text
