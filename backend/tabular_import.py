"""Reads an uploaded roster file (CSV / TSV / XLSX) into plain rows and helps match its columns to
what the voter importer needs. Pure functions, no DB access, so main.py and the tests share it.

Nothing here decides what gets saved. It only answers: what is in this file, which row looks like
the header, and which column probably holds which field. The admin confirms or corrects the guess
in the "Match columns" step, and the chosen mapping is validated again by `validate_mapping`.
"""
import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

from roster_utils import STANDARD_FIELDS

MAX_TABLE_ROWS = 50000          # data rows read after the header (same ceiling the CSV importer had)
MAX_COLS = 200                  # columns considered per row
HEAD_ROWS = 30                  # rows shown in the "raw file" preview; the header must be inside them
SAMPLE_ROWS = 5                 # data rows shown under the mapping so the admin can sanity-check it
CELL_PREVIEW_LEN = 60
MAX_XLSX_UNPACKED = 100 * 1024 * 1024   # zip-bomb guard: the upload limit is on the compressed size

TEXT_EXTS = (".csv", ".tsv", ".txt")
XLSX_EXTS = (".xlsx", ".xlsm")
SUPPORTED_HINT = "CSV, TSV, or Excel (.xlsx)"

# Targets that can take several columns joined together (e.g. First name + Surname).
MULTI_TARGETS = ("full_name", "phone")
CORE_TARGETS = (
    {"key": "student_id", "label": "Registration number", "required": True, "multi": False},
    {"key": "full_name", "label": "Full name", "required": True, "multi": True},
    {"key": "phone", "label": "Phone number", "required": False, "multi": True},
)
CORE_KEYS = tuple(t["key"] for t in CORE_TARGETS)


class TableError(ValueError):
    """A problem with the uploaded file or the chosen mapping, phrased for the admin."""


@dataclass
class Table:
    rows: list                                   # list[list[str]]; index + 1 == the row number in the file
    sheets: list = field(default_factory=list)   # workbook sheet names ([] for text files)
    sheet: str | None = None
    truncated: bool = False

    @property
    def width(self) -> int:
        return min(MAX_COLS, max((len(r) for r in self.rows[:HEAD_ROWS + SAMPLE_ROWS + 1]), default=0))

    def headers(self, header_row: int) -> list:
        if not 1 <= header_row <= len(self.rows):
            raise TableError(f"Row {header_row} is not in the file.")
        row = self.rows[header_row - 1]
        return [row[i] if i < len(row) else "" for i in range(self.width)]

    def data_rows(self, header_row: int):
        """(row_number, cells) for every non-blank row under the header."""
        for idx in range(header_row, len(self.rows)):
            cells = self.rows[idx]
            if any(cells):
                yield idx + 1, cells


# ── reading ─────────────────────────────────────────────────────────────────

def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))               # Excel stores 0772123456-as-a-number as 772123456.0
    if isinstance(v, (datetime, date)):
        return v.isoformat()[:10]
    return str(v).strip()


def _trim(cells) -> list:
    out = [_cell(c) for c in list(cells)[:MAX_COLS]]
    while out and not out[-1]:
        out.pop()
    return out


def _read_text(content: bytes, filename: str) -> Table:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")   # what Excel's "CSV" export uses on Windows
    if filename.lower().endswith(".tsv"):
        delim = "\t"
    else:
        sample = "\n".join([ln for ln in text.splitlines()[:10] if ln.strip()][:5])
        delim = max((",", ";", "\t", "|"), key=sample.count) if sample else ","
        if sample.count(delim) == 0:
            delim = ","
    limit = MAX_TABLE_ROWS + HEAD_ROWS + 1
    rows, truncated = [], False
    try:
        for rec in csv.reader(io.StringIO(text), delimiter=delim):
            if len(rows) >= limit:
                truncated = True
                break
            rows.append(_trim(rec))
    except csv.Error as e:
        raise TableError(f"That file could not be read as a table ({e}).")
    return Table(rows=rows, truncated=truncated)


def _read_xlsx(content: bytes, sheet: str | None) -> Table:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            if sum(i.file_size for i in zf.infolist()) > MAX_XLSX_UNPACKED:
                raise TableError("That workbook is too large once unpacked. Export just the voter sheet and try again.")
    except zipfile.BadZipFile:
        raise TableError("That doesn't look like a valid .xlsx file.")
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except ImportError:
        raise TableError("Excel import isn't installed on the server (pip install openpyxl). CSV files still work.")
    except Exception:
        raise TableError("That Excel file could not be opened. Try saving it again as .xlsx or CSV.")
    try:
        names = [ws.title for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"] or wb.sheetnames
        if not names:
            raise TableError("That workbook has no sheets.")
        chosen = sheet if sheet in names else names[0]
        limit = MAX_TABLE_ROWS + HEAD_ROWS + 1
        rows, truncated = [], False
        for r in wb[chosen].iter_rows(values_only=True):
            if len(rows) >= limit:
                truncated = True
                break
            rows.append(_trim(r))
        return Table(rows=rows, sheets=names, sheet=chosen, truncated=truncated)
    finally:
        wb.close()


def read_table(content: bytes, filename: str = "", sheet: str | None = None) -> Table:
    name = (filename or "").lower()
    if name.endswith(XLSX_EXTS) or (not name.endswith(TEXT_EXTS) and content[:2] == b"PK"):
        table = _read_xlsx(content, sheet)
    elif name.endswith((".xls", ".ods", ".xlsb")):
        raise TableError("Older Excel/OpenDocument formats aren't supported. In Excel choose Save As → .xlsx or CSV, then upload that.")
    else:
        table = _read_text(content, name)
    if not any(any(r) for r in table.rows):
        raise TableError("That file is empty.")
    return table


# ── guessing ────────────────────────────────────────────────────────────────

def _norm(h) -> str:
    h = re.sub(r"\(.*?\)", " ", str(h or ""))
    return re.sub(r"[^a-z0-9]+", "_", h.lower()).strip("_")


# Ordered: earlier aliases win when several columns could match. Bare "id" is last on purpose,
# because in many sheets it is a serial number rather than a registration number.
_ALIASES = {
    "student_id": ("student_id", "studentid", "registration_number", "registration_no", "reg_number", "reg_no",
                   "regno", "student_number", "student_no", "studentno", "registration", "reg", "id_no", "id"),
    "full_name": ("full_name", "fullname", "full_names", "student_name", "candidate_name", "names", "name"),
    "phone": ("phone", "phone_number", "phone_numbers", "phone_no", "mobile", "mobile_number", "mobile_no",
              "telephone", "tel", "contact", "contacts", "msisdn", "phone1", "phone_1",
              "phone2", "phone_2", "alt_phone", "alternative_phone", "secondary_phone"),
}
_FIRST = ("first_name", "firstname", "given_name", "given_names", "fname", "first")
_LAST = ("surname", "last_name", "lastname", "family_name", "lname", "other_names", "othernames", "other_name", "last")
_NAME_PARTS = frozenset(_FIRST + _LAST + ("middle_name", "middlename"))


def _field_names(f: dict) -> set:
    names = {_norm(f["key"]), _norm(f.get("label"))}
    if f.get("standard"):
        std = next((s for s in STANDARD_FIELDS if s["key"] == f["key"]), None)
        names |= {_norm(a) for a in (std or {}).get("aliases", ())}
    return names - {""}


def targets_for(fields) -> list:
    """Everything the admin can map: the three core columns plus the org's enabled voter fields."""
    out = [dict(t) for t in CORE_TARGETS]
    out += [{"key": f["key"], "label": f.get("label") or f["key"], "required": False, "multi": False}
            for f in (fields or []) if f.get("enabled")]
    return out


def guess_header_row(rows, fields=None) -> int:
    known = set().union(*_ALIASES.values(), _FIRST, _LAST)
    for f in fields or []:
        if f.get("enabled"):
            known |= _field_names(f)
    best_row, best_score = 0, 0
    for i, row in enumerate(rows[:15]):
        score = sum(1 for c in row if _norm(c) in known)
        if score > best_score:
            best_row, best_score = i, score
    if best_score:
        return best_row + 1
    return next((i + 1 for i, r in enumerate(rows[:HEAD_ROWS]) if any(r)), 1)


def suggest_mapping(headers, fields=None) -> dict:
    """{target: [column index, ...]} — a starting point for the admin, never applied unseen."""
    norm = [_norm(h) for h in headers]
    used: set = set()
    mapping: dict = {}

    def find(names):
        for n in names:
            for i, h in enumerate(norm):
                if h == n and i not in used:
                    return i
        return None

    sid = find(_ALIASES["student_id"])
    if sid is not None:
        mapping["student_id"] = [sid]
        used.add(sid)
    name = find(_ALIASES["full_name"])
    if name is not None:
        mapping["full_name"] = [name]
        used.add(name)
    else:
        # "First name" + "Surname", or "Surname" + "Other names": join them in the order they appear.
        parts = [i for i, h in enumerate(norm) if h in _NAME_PARTS and i not in used][:3]
        if len(parts) >= 2:
            mapping["full_name"] = parts
            used.update(parts)
    phones = []
    for _ in range(3):
        p = find(_ALIASES["phone"])
        if p is None:
            break
        phones.append(p)
        used.add(p)
    if phones:
        mapping["phone"] = sorted(phones)
    for f in fields or []:
        if f.get("enabled"):
            hit = find(sorted(_field_names(f)))
            if hit is not None:
                mapping[f["key"]] = [hit]
                used.add(hit)
    return mapping


def validate_mapping(raw, allowed_attr_keys) -> dict:
    """Checks a client-supplied mapping and returns it as {target: [int, ...]}."""
    if not isinstance(raw, dict):
        raise TableError("The column mapping is not valid.")
    allowed = set(CORE_KEYS) | set(allowed_attr_keys)
    out, seen = {}, {}
    for key, cols in raw.items():
        if key not in allowed:
            raise TableError(f"'{key}' is not a field that can be imported.")
        cols = cols if isinstance(cols, list) else [cols]
        idxs = []
        for c in cols:
            if c in (None, ""):
                continue
            if isinstance(c, bool) or not isinstance(c, (int, str)) or not str(c).lstrip("-").isdigit():
                raise TableError("The column mapping is not valid.")
            i = int(c)
            if not 0 <= i < MAX_COLS:
                raise TableError("The column mapping points outside the file.")
            if i in seen and seen[i] != key:
                raise TableError(f"One column can't feed both '{seen[i]}' and '{key}'.")
            if i in idxs:
                continue
            seen[i] = key
            idxs.append(i)
        if len(idxs) > 1 and key not in MULTI_TARGETS:
            raise TableError(f"'{key}' can only come from one column.")
        if idxs:
            out[key] = idxs
    missing = [t["label"] for t in CORE_TARGETS if t["required"] and t["key"] not in out]
    if missing:
        raise TableError("Choose a column for: " + ", ".join(missing) + ".")
    return out


def _preview_cells(cells, width) -> list:
    return [(cells[i][:CELL_PREVIEW_LEN] if i < len(cells) else "") for i in range(width)]


def describe_table(table: Table, header_row: int | None, fields) -> dict:
    """Everything the "Match columns" screen needs, in one response."""
    if header_row is None:
        header_row = guess_header_row(table.rows, fields)
    if not 1 <= header_row <= min(len(table.rows), HEAD_ROWS):
        raise TableError(f"The header row must be one of the first {min(len(table.rows), HEAD_ROWS)} rows of the file.")
    headers = table.headers(header_row)
    width = table.width
    sample = []
    data_count = 0
    for row_num, cells in table.data_rows(header_row):
        data_count += 1
        if len(sample) < SAMPLE_ROWS:
            sample.append({"row": row_num, "cells": _preview_cells(cells, width)})
    return {
        "sheets": table.sheets, "sheet": table.sheet,
        "header_row": header_row, "width": width,
        "headers": [h[:CELL_PREVIEW_LEN] for h in headers],
        "raw_preview": [{"row": i + 1, "cells": _preview_cells(r, width)} for i, r in enumerate(table.rows[:HEAD_ROWS])],
        "sample": sample, "data_rows": data_count, "truncated": table.truncated,
        "targets": targets_for(fields),
        "suggested_mapping": suggest_mapping(headers, fields),
    }
